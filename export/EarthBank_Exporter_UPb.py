#/ Type: Exporter
#/ Name: EarthBank U-Pb Exporter
#/ Authors: Bence Paul and Joe Petrus
#/ Description: An export script for creating an EarthBank formatted spreadsheet for importing into the EarthBank database.
#/ References: None
#/ Version: 0.1
#/ Contact: iolitesupport@icpms.com

import os
import subprocess
from shutil import copy2
from pprint import pprint

from openpyxl import load_workbook

from iolite.Qt import Qt
from iolite.QtCore import QFile, QIODevice, QEventLoop
from iolite.QtGui import QComboBox, QHeaderView, QLabel, QLineEdit, QMessageBox, QStackedWidget, QTableWidget, QTableWidgetItem, QToolButton, QVBoxLayout, QWidget, QDoubleSpinBox
from iolite.QtGui import QCheckBox
from iolite.QtUiTools import QUiLoader

'''
Paths: tell the script where to find the template file and the UI file. These are hardcoded for now, but could be made more flexible in the future.
'''
TEMPLATE_PATH = '/Users/bence/Dropbox/testing/AGN/UPb_Template/UPbDatapoint_LAICPMS.template.v2026-06-22.xlsx'
UI_FILE_PATH = '/Users/bence/iolite4-python-examples/export/EarthBank_Exporter_UPb.ui'

'''
Variables for your lab.
'''
USERS = ['Alan Grieg', 'Ashlea Wainwright', 'Bence Paul', 'Brandon Mahan', 'Janet Hergt', 'Jon Woodhead', 'Roland Maas']
LAB = 'Isotope Geochemistry,  The University of Melbourne'

LASER_WAVELENGTHS = ['193 nm', '213 nm', '266 nm']
LASER_PULSE_WIDTHS = ['< 5 ns', '5-10 ns', '> 10 ns']
CELL_MODEL = ['TV2', 'TV3', 'Helex']
CARRIER_GAS = ['He', 'Ar', 'N2']
MIXING_DEVICES = ['Squid', 'ESL Mixing bulb', 'Other']

MULTICOLLECTOR_MACHINE_NAMES = ['Thermo Nepture', 'Neoma', 'Nu Plasma II/III']
TOF_MACHINE_NAMES = ['Vitesse Text', 'icpTOF']

class SettingsWidget(QWidget):

    def __init__(self, parent=None):
        super().__init__(parent)

        self.setAttribute(Qt.WA_DeleteOnClose, True)

        self.setLayout(QVBoxLayout())

        # Load the template file for the AGN export. This is an Excel file that contains the correct formatting for the AGN database.
        self.export_template = QFile(TEMPLATE_PATH)
        if not self.export_template.exists():
            raise RuntimeError(f'Could not find export template {self.export_template.fileName()}')

        if not self.export_template.open(QIODevice.ReadOnly):
            raise RuntimeError(f'Could not load export template file for reading: {self.export_template.fileName()}')

        # Load the UI file
        self.ui_file = QFile(UI_FILE_PATH)

        if not self.ui_file.exists():
            raise RuntimeError(f'Could not find settings ui {self.ui_file.fileName()}')

        if not self.ui_file.open(QIODevice.ReadOnly):
            raise RuntimeError(f'Could not load settings ui file for reading: {self.ui_file.fileName()}')

        ui = QUiLoader().load(self.ui_file, self)
        self.layout().addWidget(ui)
        self.layout().setContentsMargins(0, 0, 0, 0)

        self.stackedWidget = ui.findChild(QStackedWidget, 'stackedWidget')
        self.cancelButton = ui.findChild(QToolButton, 'cancelButton')
        self.continueButton = ui.findChild(QToolButton, 'continueButton')
        self.backButton = ui.findChild(QToolButton, 'backButton')
        self.stepLabel = ui.findChild(QLabel, 'stepLabel')

        self.userComboBox = ui.findChild(QComboBox, 'userComboBox')
        self.labLineEdit = ui.findChild(QLineEdit, 'labLineEdit')
        self.litLineEdit = ui.findChild(QLineEdit, 'litLineEdit')
        self.fundingLineEdit = ui.findChild(QLineEdit, 'fundingLineEdit')
        self.scaleComboBox = ui.findChild(QComboBox, 'scaleComboBox')
        self.techniqueComboBox = ui.findChild(QComboBox, 'techniqueComboBox')
        self.uncertTypeComboBox = ui.findChild(QComboBox, 'uncertaintyTypeComboBox')
        self.trimSigFigsCheckBox = ui.findChild(QCheckBox, 'trimSigFigsCheckBox')
        self.sessionIDComboBox = ui.findChild(QComboBox, 'sessionIDComboBox')
        self.groupsTable = ui.findChild(QTableWidget, 'groupsTableWidget')
        self.duplicatesLabel = ui.findChild(QLabel, 'ratiosLabel')
        self.ratiosTable = ui.findChild(QTableWidget, 'ratiosTableWidget')
        self.laserFluenceDoubleSpinBox = ui.findChild(QDoubleSpinBox, 'laserFluenceDoubleSpinBox')
        self.laserWavelengthComboBox = ui.findChild(QComboBox, 'laserWavelengthComboBox')
        self.laserPulseWidthComboBox = ui.findChild(QComboBox, 'laserPulseWidthComboBox')
        self.cellModelComboBox = ui.findChild(QComboBox, 'cellModelComboBox')
        self.carrierGasComboBox = ui.findChild(QComboBox, 'carrierGasComboBox')
        self.carrierGasFlowRateSpinBox = ui.findChild(QDoubleSpinBox, 'carrierGasFlowRateDoubleSpinBox')
        self.mixingDeviceComboBox = ui.findChild(QComboBox, 'mixingDeviceComboBox')

        # Get the valid options from the template file
        template_wb = load_workbook(TEMPLATE_PATH)
        lookup_tables_ws = template_wb['Lookup Tables']
        lookup_values = {}
        for column in range(1, lookup_tables_ws.max_column + 1):
            lookup_name = lookup_tables_ws.cell(row=2, column=column).value
            if lookup_name in (None, 'Description'):
                continue

            values = []
            for row in range(3, lookup_tables_ws.max_row + 1):
                value = lookup_tables_ws.cell(row=row, column=column).value
                if value in (None, ''):
                    break
                values.append(value)
            lookup_values[lookup_name] = values

        # Some sanity checks to make sure that the template file contains the expected values
        if 'U-Pb Analytical Technique' not in lookup_values:
            raise RuntimeError('Could not find "U-Pb Analytical Technique" in the template file. Please check that the template file is correct.')

        if 'Mineral Type' not in lookup_values:
            raise RuntimeError('Could not find "Mineral Type" in the template file. Please check that the template file is correct.')

        # Fill mineral type combo box
        self.mineralTypeComboBox = ui.findChild(QComboBox, 'mineralComboBox')
        self.mineralTypeComboBox.addItems(lookup_values['Mineral Type'])
        self.mineralTypeComboBox.setCurrentText('Zircon')
        
        self.techniqueComboBox.addItems(lookup_values['U-Pb Analytical Technique'])
        self.techniqueComboBox.setCurrentText('LA-ICP-MS')
        
        self.userComboBox.addItems(USERS)
        self.labLineEdit.setText(LAB)
        self.uncertTypeComboBox.addItems(['2 standard error', '1σ', '2σ', '95%'])
        self.sessionIDComboBox.addItems([data.sessionUUID(), 'File Path'])

        self.laserFluenceDoubleSpinBox.setValue(1.0)
        self.laserWavelengthComboBox.addItems(LASER_WAVELENGTHS)
        self.laserWavelengthComboBox.setCurrentText('193 nm')
        self.laserPulseWidthComboBox.addItems(LASER_PULSE_WIDTHS)
        self.laserPulseWidthComboBox.setCurrentText('< 5 ns')
        self.cellModelComboBox.addItems(CELL_MODEL)
        self.cellModelComboBox.setCurrentText('TV2')
        self.carrierGasComboBox.addItems(CARRIER_GAS)
        self.carrierGasComboBox.setCurrentText('He')
        self.mixingDeviceComboBox.addItems(MIXING_DEVICES)
        self.mixingDeviceComboBox.setCurrentText('Squid')
        self.carrierGasFlowRateSpinBox.setValue(1.0)

        '''
        Get primary RMs (calibrants) as these should not be included in exported results.
        '''
        calibrants = []

        for ch in data.timeSeriesList(data.Input):
            extStd = ch.property('External standard')
            if extStd is not None:
                calibrants.append(extStd)

        # Set up the groups table
        self.groupsTable.setColumnCount(3)
        self.groupsTable.setHorizontalHeaderLabels(['Group Name', 'Type', 'Selection Count'])
        self.groupsTable.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.groupsTable.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        self.groupsTable.setEditTriggers(QTableWidget.NoEditTriggers)
        def createGroupTableEntry(group, type):
            row = self.groupsTable.rowCount
            self.groupsTable.insertRow(row)

            groupNameItem = QTableWidgetItem(group.name)
            if group.name in calibrants:
                groupNameItem.setCheckState(Qt.Unchecked)
            else:
                groupNameItem.setCheckState(Qt.Checked)
            self.groupsTable.setItem(row, 0, groupNameItem)

            groupTypeItem = QTableWidgetItem(type)
            self.groupsTable.setItem(row, 1, groupTypeItem)

            selCountItem = QTableWidgetItem(str(group.count))
            selCountItem.setTextAlignment(Qt.AlignHCenter)
            self.groupsTable.setItem(row, 2, selCountItem)


        for group in data.selectionGroupList(data.ReferenceMaterial):
            createGroupTableEntry(group, 'Reference Material')

        for group in data.selectionGroupList(data.Sample):
            createGroupTableEntry(group, 'Sample')

        # Set up the ratios table
        self.ratiosTable.setColumnCount(2)
        self.ratiosTable.setHorizontalHeaderLabels(['Result', 'Channel'])
        self.ratiosTable.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.ratiosTable.horizontalHeader().setSectionResizeMode(1, QHeaderView.Stretch)
        self.ratiosTable.verticalHeader().hide() 
        # self.ratiosTable.setEditTriggers(QTableWidget.NoEditTriggers)

        U_Pb_field_defaults = {
            'U concentration (µg.g-1)': 'Approx_U_PPM',
            'Th concentration (µg.g-1)': 'Approx_Th_PPM',
            '206Pb intensity (mV or cps)': 'Pb206_CPS',
            '204Pb intensity (mV or cps)': 'Pb204_CPS',
            'Total Pb concentration (µg.g-1)': 'Approx_Pb_PPM',
            '206Pb/238U Ratio': 'Final Pb206/U238',
            '238U/206Pb Ratio': 'Final U238/Pb206',
            '207Pb/206Pb Ratio': 'Final Pb207/Pb206',
            '207Pb/235U Ratio': 'Final Pb207/U235',
            '208Pb/232Th Ratio': 'Final Pb208/Th232',
            '206Pb/238U Date (Ma)': 'Final Pb206/U238 age',
            '207Pb/206Pb Date (Ma)': 'Final Pb207/Pb206 age', 
            '207Pb/235U Date (Ma)': 'Final Pb207/U235 age', 
            '208Pb/232Th Date (Ma)': 'Final Pb208/Th232 age',
        }

        channels = [ch.name for ch in data.timeSeriesList()]
        channels.insert(0, 'None')

        for option, default_channel in U_Pb_field_defaults.items():
            row = self.ratiosTable.rowCount
            self.ratiosTable.insertRow(row)
            self.ratiosTable.setItem(row, 0, QTableWidgetItem(option))
            chItem = QComboBox()
            chItem.addItems(channels)
            chItem.setCurrentText(default_channel if default_channel in channels else 'None')
            self.ratiosTable.setCellWidget(row, 1, chItem)

        self.backButton.clicked.connect(self.previousTab)
        self.continueButton.clicked.connect(self.nextTab)
        self.cancelButton.clicked.connect(lambda: self.close())

        self.stackedWidget.setCurrentIndex(0)
        self.stepLabel.setText(f'Step {self.stackedWidget.currentIndex+1} of 4')

    def previousTab(self):
        current_index = self.stackedWidget.currentIndex
        self.stackedWidget.setCurrentIndex(current_index-1)
        self.stepLabel.setText(f'Step {self.stackedWidget.currentIndex+1} of 4')
        self.continueButton.setText("Continue")

    def nextTab(self):
        current_index = self.stackedWidget.currentIndex
        # Check that at least one group is selected in groupsTable
        if current_index == 2 and all(self.groupsTable.item(row, 0).checkState() != Qt.Checked for row in range(self.groupsTable.rowCount)):
            QMessageBox.warning(self, 'No groups selected', 'Please select at least one group before continuing.')
            return
        # Check that not all groups are selected in groupsTable (the user should leave at least one group unselected... most likely)
        if current_index == 2 and all(self.groupsTable.item(row, 0).checkState() == Qt.Checked for row in range(self.groupsTable.rowCount)):
            response = QMessageBox.question(
                self,
                'All groups selected',
                'All groups are selected. Click OK to continue, or Cancel to deselect one of the groups.',
                QMessageBox.Ok | QMessageBox.Cancel,
                QMessageBox.Cancel,
            )
            if response == QMessageBox.Cancel:
                return

        # If we're already on the last tab, export the data
        if current_index == self.stackedWidget.count - 1:
            self.exportData()
            print('Export complete.')
            return
        self.stackedWidget.setCurrentIndex(current_index+1)
        self.stepLabel.setText(f'Step {self.stackedWidget.currentIndex+1} of 4')
        if self.stackedWidget.currentIndex == self.stackedWidget.count - 1:
            self.continueButton.setText("Export Data")
        else:
            self.continueButton.setText("Continue")


    def exportData(self):
        print('Exporting data...')

        try:
            # NOTE: the export_filepath variable is automatically set for the script by iolite, 
            # and is the path to the file that the user has selected to export to.
            fp = copy2(TEMPLATE_PATH, export_filepath)
        except FileNotFoundError:
            print('Could not find template file')
            return
        except PermissionError:
            print('Permission denied to write to export file')
            return
        except:
            print('Could not create copy of template file')
            return

        wb = load_workbook(fp)

        try:
            dps_ws = wb['UPb Datapoints']
            upb_ws = wb['UPbSpotData']
            icpms_ws = wb['ICPMS']
        except KeyError:
            print('Could not find expected sheet in template file')
            return

        # Get the column indices for the various data points in the template file
        dps_col_indices = {}
        for col in dps_ws.iter_cols(min_row=3, max_row=3):
            header = col[0].value
            if header:
                dps_col_indices[header] = col[0].column

        # Repeat for UPbSpotData sheet
        upb_col_indices = {}
        for col in upb_ws.iter_cols(min_row=4, max_row=4):
            header = col[0].value
            if header:
                upb_col_indices[header] = col[0].column

        # Repeat for ICP-MS sheet
        icpms_col_indices = {}
        for col in icpms_ws.iter_cols(min_row=4, max_row=4):
            header = col[0].value
            if header:
                icpms_col_indices[header] = col[0].column

        # Get list of groups to export and their type
        groups_to_export = []
        for row in range(self.groupsTable.rowCount):
            grpName = self.groupsTable.item(row, 0).text()
            if self.groupsTable.item(row, 0).checkState() == Qt.Checked:
                groups_to_export.append(grpName)

        groupCounter = 5
        rowNo = 5
        for g in groups_to_export:
            print(f'Exporting group: {g}')
            group = data.selectionGroup(g)
            datapoint_name = group.name
            dps_ws.cell(groupCounter, 1, value=datapoint_name)
            if group.type == data.Sample:
                dps_ws.cell(groupCounter, 2, value=group.name)
            elif group.type == data.ReferenceMaterial:
                dps_ws.cell(groupCounter, 3, value=group.name)
            
            # Export metadata to the template file
            dps_ws.cell(groupCounter, dps_col_indices['U-Pb Analytical Technique'], value = self.techniqueComboBox.currentText)
            dps_ws.cell(groupCounter, 6, value = data.sessionFilePath())
            dps_ws.cell(groupCounter, dps_col_indices['Analyst'], value = self.userComboBox.currentText)
            dps_ws.cell(groupCounter, dps_col_indices['Laboratory'], value = self.labLineEdit.text)
            dps_ws.cell(groupCounter, dps_col_indices['Mineral Type'], value = self.mineralTypeComboBox.currentText)
            # Note the typo in 'Associated Litterature' which is in the template. When this gets changed, we'll need to update this line
            dps_ws.cell(groupCounter, dps_col_indices['Associated litterature'], value = self.litLineEdit.text)
            dps_ws.cell(groupCounter, dps_col_indices['Funding'], value = self.fundingLineEdit.text)

            # Now export U-Pb data
            if rowNo == 5:
                rowNo += 1 # Skip the header row for U-Pb data

            # Keep track of repetition rates, spot heights, and spot widths for each selection
            rep_rates = set()
            spot_heights = set()
            spot_widths = set()
            
            for sel in group.selections():
                upb_ws.cell(row=rowNo, column=1, value=datapoint_name)
                upb_ws.cell(row=rowNo, column=3, value=sel.name)
                upb_ws.cell(row=rowNo, column=4, value=sel.startTime.toString('yyyy-MM-dd HH:mm:ss'))
                upb_ws.cell(row=rowNo, column=6, value=sel.comment)

                rep_rates.add(sel.property('Rep Rate'))
                spot_heights.add(sel.property('Spot Height'))
                spot_widths.add(sel.property('Spot Width'))

                for row in range(self.ratiosTable.rowCount):
                    field_name = self.ratiosTable.item(row, 0).text() if self.ratiosTable.item(row, 0) is not None else None
                    ch_name = self.ratiosTable.cellWidget(row, 1).currentText if self.ratiosTable.cellWidget(row, 1) is not None else None

                    if field_name is None or ch_name is None or ch_name == 'None':
                        #print(f'Missing field name or channel name at row {row}: field_name={field_name}, ch_name={ch_name}')
                        continue

                    result = data.result(sel, data.timeSeries(ch_name)).value()
                    uncert = data.result(sel, data.timeSeries(ch_name)).propagatedUncertainty()

                    upb_ws.cell(row=rowNo, column=upb_col_indices[field_name], value=result)
                    if isinstance(uncert, float) and (uncert != uncert):  # Check for NaN
                        pass
                        #print(f'No propagated uncertainty for field {field_name} at row {row}')
                    else:
                        upb_ws.cell(row=rowNo, column=upb_col_indices[field_name] + 1, value=uncert)

                rowNo += 1

            # If any of the rep rate, spot height, or spot width have a single unique value, export them to the ICP-MS sheet
            if len(rep_rates) == 1 or len(spot_heights) == 1 or len(spot_widths) == 1:
                icpms_ws.cell(row=groupCounter, column=icpms_col_indices['datapointName'], value=datapoint_name)
            if len(rep_rates) == 1:
                icpms_ws.cell(row=groupCounter, column=icpms_col_indices['laserMetadata_laserRepetitionRate'], value=rep_rates.pop())
            else:
                print(f'Not exporting rep rates for group {group.name} because there are multiple rep rates reported: {rep_rates}')
            if len(spot_heights) == 1 and len(spot_widths) == 1:
                spot_height = next(iter(spot_heights))
                spot_width = next(iter(spot_widths))

                if spot_height == spot_width:
                    icpms_ws.cell(
                        row=groupCounter,
                        column=icpms_col_indices['laserMetadata_laserSpotSize'],
                        value=spot_height
                    )
            else:
                print(f'Not exporting spot size for group {group.name} because there are multiple spot heights reported: {spot_heights} or multiple spot widths reported: {spot_widths}')

            # Export other laser metadata
            icpms_ws.cell(row=groupCounter, column=icpms_col_indices['laserMetadata_laserFluence'], value=self.laserFluenceDoubleSpinBox.value)
            # icpms_ws.cell(row=groupCounter, column=icpms_col_indices['laserMetadata_laserWavelength'], value=self.laserWavelengthComboBox.currentText())
            icpms_ws.cell(row=groupCounter, column=icpms_col_indices['laserMetadata_pulseWidthValue'], value=self.laserPulseWidthComboBox.currentText)
            icpms_ws.cell(row=groupCounter, column=icpms_col_indices['laserMetadata_cellModel'], value=self.cellModelComboBox.currentText)
            icpms_ws.cell(row=groupCounter, column=icpms_col_indices['laserMetadata_carrierGas'], value=self.carrierGasComboBox.currentText)
            icpms_ws.cell(row=groupCounter, column=icpms_col_indices['laserMetadata_mixingDevice'], value=self.mixingDeviceComboBox.currentText)
            icpms_ws.cell(row=groupCounter, column=icpms_col_indices['laserMetadata_carrierGasFlow'], value=self.carrierGasFlowRateSpinBox.value)

            groupCounter += 1

        wb.save(fp)
        # Now use the system open the file in Excel
        subprocess.run(['open', export_filepath])
        self.close()

widget = SettingsWidget()
widget.show()
loop = QEventLoop()
widget.destroyed.connect(lambda: loop.quit)
loop.exec() # wait ...
print('finished')
